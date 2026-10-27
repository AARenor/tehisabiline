#!/usr/bin/env python3
"""Build assets/downloads/nis2-auditi-toendite-mall.xlsx (NIS2 audit evidence register).

Source of truth for the downloadable evidence template served by /nis2/audit/.
Rows map the ten minimum measures of NIS2 (EU) 2022/2555 art 21(2)(a)-(j).

Run from the repo root:  python3 tools/build-nis2-evidence-xlsx.py
Requires openpyxl (system python3-openpyxl).
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = Path(__file__).resolve().parent.parent / "assets" / "downloads" / "nis2-auditi-toendite-mall.xlsx"

INK = "FF1F3B4D"
MUTED = "FF5A6B75"
BAND = "FFF2F6F8"
LINE = Side(style="thin", color="FFD3DDE3")

ROWS = [
    ("1", "21(2)(a)", "Riskianalüüs ja turvapoliitika", "Kinnitatud riskiregister ja turvapoliitika", "Küberturbe juht", "Vali 3 riski ja kontrolli ravi ning jääkriski"),
    ("2", "21(2)(b)", "Intsidentide käsitlemine", "Intsidendiprotseduur ja kontaktid", "IT-juht", "Simuleeri intsident ja kontrolli 24/72 tunni rada"),
    ("3", "21(2)(c)", "Toimepidevus ja varundus", "Varundusplaan ja taastetesti raport", "IT-juht", "Kontrolli viimast taastetesti ja RPO/RTO vastavust"),
    ("4", "21(2)(d)", "Tarneahela turve", "Kriitiliste tarnijate nimekiri ja lepingud", "Hanked", "Vali 3 tarnijat ja kontrolli due diligence't"),
    ("5", "21(2)(e)", "Turvaline arendus ja haavatavused", "Haavatavuste protsess ja parandusajad", "Arendusjuht", "Kontrolli 5 haavatavuse parandamise tõendit"),
    ("6", "21(2)(f)", "Tõhususe testimine", "Testiplaan ja viimased tulemused", "Turvajuht", "Kontrolli, kas testid katavad kriitilised süsteemid"),
    ("7", "21(2)(g)", "Küberhügieen ja koolitus", "Koolituskava ja osalemise tõendid", "Personal", "Kontrolli koolituse katvust ja värskust"),
    ("8", "21(2)(h)", "Krüpteerimine", "Krüpteerimise poliitika ja kasutuspõhimõtted", "IT-juht", "Kontrolli andmete liikumist ja krüpteerimist"),
    ("9", "21(2)(i)", "Ligipääsukontroll", "Privileegide maatriks ja ülevaatuse tõendid", "IT-juht", "Vali 5 kontot ja kontrolli õigusi"),
    ("10", "21(2)(j)", "MFA ja turvaline side", "MFA katvus ja erisused", "IT-juht", "Kontrolli kriitiliste kontode MFA-d"),
]

HEADERS = [
    ("Meede", 7),
    ("Säte", 11),
    ("Nõue", 30),
    ("Tõend, mida koguda", 34),
    ("Omanik", 15),
    ("Testi kirjeldus", 34),
    ("Testi kuupäev", 14),
    ("Tulemus", 17),
    ("Puudus (leid)", 30),
    ("Parandus (tegevus)", 30),
    ("Tähtaeg", 12),
    ("Staatus", 16),
    ("Tõendi asukoht (link või tee)", 30),
]

RESULTS = ["Vastab", "Osaliselt vastab", "Ei vasta", "Pole hinnatud"]
STATUSES = ["Pole alustatud", "Töös", "Valmis", "Edasi lükatud"]
LAST_ROW = 200  # validation/formatting reach further than the seeded rows


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tõendid"
    ws.sheet_properties.tabColor = INK[2:]

    for col, (name, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="center" if col in (1, 2, 7, 11) else "left")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 34

    for r, (measure, article, requirement, evidence, owner, test) in enumerate(ROWS, start=2):
        values = [measure, article, requirement, evidence, owner, test, None, "Pole hinnatud",
                  None, None, None, "Pole alustatud", None]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if c in (1, 2, 7, 8, 11, 12) else "left")
            cell.border = Border(bottom=LINE)
            if c in (7, 11):
                cell.number_format = "DD.MM.YYYY"
            if r % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=BAND)
        ws.row_dimensions[r].height = 34

    # Status and result are pick-lists: free text would make the register unsortable.
    for status_list, col_letter in ((RESULTS, "H"), (STATUSES, "L")):
        dv = DataValidation(type="list", formula1='"' + ",".join(status_list) + '"',
                            allow_blank=True, showDropDown=False)
        dv.errorTitle = "Vali loendist"
        dv.error = "Kasuta ühte lubatud väärtust: " + ", ".join(status_list)
        dv.promptTitle = "Lubatud väärtused"
        dv.prompt = ", ".join(status_list)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{LAST_ROW}")

    fills = {
        "Vastab": ("FFE3F5E9", "FF1E6B3A"),
        "Osaliselt vastab": ("FFFFF2D0", "FF7A5300"),
        "Ei vasta": ("FFFCE4E4", "FF9B1C1C"),
        "Pole hinnatud": ("FFEDF1F4", "FF5A6B75"),
    }
    for value, (bg, fg) in fills.items():
        ws.conditional_formatting.add(
            f"H2:H{LAST_ROW}",
            FormulaRule(formula=[f'$H2="{value}"'],
                        fill=PatternFill("solid", fgColor=bg), font=Font(color=fg)),
        )
    for value, (bg, fg) in {
        "Valmis": ("FFE3F5E9", "FF1E6B3A"),
        "Töös": ("FFE4EDF7", "FF1B4E7A"),
        "Pole alustatud": ("FFEDF1F4", "FF5A6B75"),
        "Edasi lükatud": ("FFFFF2D0", "FF7A5300"),
    }.items():
        ws.conditional_formatting.add(
            f"L2:L{LAST_ROW}",
            FormulaRule(formula=[f'$L2="{value}"'],
                        fill=PatternFill("solid", fgColor=bg), font=Font(color=fg)),
        )
    # Overdue action, unless the row is already closed.
    ws.conditional_formatting.add(
        f"K2:K{LAST_ROW}",
        FormulaRule(formula=[f'AND($K2<>"",$L2<>"Valmis",ISNUMBER($K2),$K2<TODAY())'],
                    fill=PatternFill("solid", fgColor="FFFCE4E4"), font=Font(color="FF9B1C1C", bold=True)),
    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{len(ROWS) + 1}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    guide = wb.create_sheet("Juhend")
    guide.sheet_properties.tabColor = MUTED[2:]
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["B"].width = 74
    guide.column_dimensions["C"].width = 18

    def block(row, text, size=11, bold=False, color="FF000000", italic=False):
        cell = guide.cell(row=row, column=1, value=text)
        cell.font = Font(size=size, bold=bold, color=color, italic=italic)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        return cell

    block(1, "NIS2 auditi tõendite mall", size=16, bold=True, color=INK)
    block(2, "Töömall tõendite kogumiseks NIS2 direktiivi (EL) 2022/2555 artikli 21 lõike 2 kümne miinimumnõude kohta. Täida lehel „Tõendid“ üks rida iga meetme kohta.", color=MUTED)
    guide.merge_cells("A2:C2")
    guide.row_dimensions[2].height = 32

    block(4, "Kuidas täita", size=12, bold=True, color=INK)
    steps = [
        "Määra iga meetme omanik ja kirjuta testi kirjeldus nii, et teine inimene saaks sama testi korrata.",
        "Kirjuta testi kuupäev ja tulemus rippmenüüst. Tulemus peab põhinema nähtud tõendil, mitte vastusel küsimusele „kas see on olemas?“.",
        "Pane tõendi asukoht (jagatud kausta tee, pileti number või link) veergu „Tõendi asukoht“. Ilma asukohata ei ole tõend auditiks kasutatav.",
        "Iga puuduse kohta kirjuta parandus ja tähtaeg. Möödunud tähtaeg värvub punaseks, kuni staatus on „Valmis“.",
        "Enne auditi algust vaata juhatusele üle kokkuvõte lehe „Juhend“ allosas ja otsusta, millised puudused viiakse tegevuskavasse.",
    ]
    for i, text in enumerate(steps, start=5):
        block(i, f"{i - 4}.", bold=True)
        cell = guide.cell(row=i, column=2, value=text)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[i].height = 30

    block(11, "Veergude tähendus", size=12, bold=True, color=INK)
    legend = [
        ("Meede / Säte", "NIS2 artikli 21 lõike 2 punkt; hoiab rea seotud direktiivi tekstiga."),
        ("Nõue", "Meetme nimi sellisel kujul, nagu seda kohustuslase dokumentatsioonis kasutatakse."),
        ("Tõend, mida koguda", "Dokument või väljavõte, mida auditeerija küsib."),
        ("Tulemus", "Vastab / Osaliselt vastab / Ei vasta / Pole hinnatud."),
        ("Puudus", "Konkreetne leid: mis puudub või ei toimi, mitte üldine hinnang."),
        ("Staatus", "Pole alustatud / Töös / Valmis / Edasi lükatud."),
    ]
    for i, (a, b) in enumerate(legend, start=12):
        cell = guide.cell(row=i, column=1, value=a)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell2 = guide.cell(row=i, column=2, value=b)
        cell2.alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[i].height = 26

    block(19, "Kokkuvõte (arvutatakse lehelt „Tõendid“)", size=12, bold=True, color=INK)
    sheet = "'Tõendid'"
    col_c = f"{sheet}!$C$2:$C${LAST_ROW}"
    col_h = f"{sheet}!$H$2:$H${LAST_ROW}"
    col_k = f"{sheet}!$K$2:$K${LAST_ROW}"
    col_l = f"{sheet}!$L$2:$L${LAST_ROW}"
    summary = [
        ("Ridu kokku", f"=COUNTA({col_c})"),
        ("Valmis", f'=COUNTIF({col_l},"Valmis")'),
        ("Töös", f'=COUNTIF({col_l},"Töös")'),
        ("Pole alustatud", f'=COUNTIF({col_l},"Pole alustatud")'),
        ("Edasi lükatud", f'=COUNTIF({col_l},"Edasi lükatud")'),
        ("Tulemus: vastab", f'=COUNTIF({col_h},"Vastab")'),
        ("Tulemus: osaliselt vastab", f'=COUNTIF({col_h},"Osaliselt vastab")'),
        ("Tulemus: ei vasta", f'=COUNTIF({col_h},"Ei vasta")'),
        ("Üle tähtaja", f'=SUMPRODUCT(({col_k}<>"")*({col_k}<TODAY())*({col_l}<>"Valmis"))'),
    ]
    for i, (label, formula) in enumerate(summary, start=20):
        a = guide.cell(row=i, column=1, value=label)
        a.font = Font(size=11)
        b = guide.cell(row=i, column=3, value=formula)
        b.font = Font(size=11, bold=True)
        b.alignment = Alignment(horizontal="center")

    block(31, "Mida see fail ei ole", size=12, bold=True, color=INK)
    caveat = guide.cell(row=32, column=1, value=(
        "See on töövahend, mitte auditiraport, sertifikaat ega vastavusdeklaratsioon. NIS2 ei nõua ühte sertifikaati: "
        "auditi järeldus peab põhinema ulatusel, kriteeriumidel ja nähtud tõenditel. Täidetud tabel ilma tõendite asukohtadeta "
        "ei ole auditiks piisav. Kontrolli enne kasutamist kehtivat NIS2/Eesti KüTS redaktsiooni."
    ))
    caveat.alignment = Alignment(vertical="top", wrap_text=True)
    guide.merge_cells("A32:C32")
    guide.row_dimensions[32].height = 58

    block(34, "Allikas: NIS2 direktiiv (EL) 2022/2555, art 21 lg 2 · tehisabiline.ee/nis2/audit/", size=9, color=MUTED, italic=True)

    wb.properties.title = "NIS2 auditi tõendite mall"
    wb.properties.subject = "NIS2 direktiivi art 21 lg 2 tõendite register"
    wb.properties.creator = "tehisabiline.ee"
    wb.properties.keywords = "NIS2, audit, tõendid, art 21, KüTS"
    wb.properties.description = "Töömall NIS2 auditi tõendite kogumiseks. Allikas: https://tehisabiline.ee/nis2/audit/"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(OUT.parent.parent.parent)} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
